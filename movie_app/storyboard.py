"""Storyboards (콘티) in Excel and as pictures.

- read(): a 콘티 workbook whatever its columns: the header row is found by its words (장면/나레이션/화면설명/
  소스/템플릿/길이 ... in several spellings), or the columns are named.
- script(): the narration to make the voice from, a paragraph a row (a row without narration a pause).
- build(): a project from the rows: a picture or clip from the source column, a graphic scene from a template,
  a title card from on-screen words; timed to a voice made from the script when there is one.
- export(): a project back as a 콘티 workbook (thumbnails, times, narration), with a source list and a
  narration check; compare(): the narration written against what the voice says; contact_sheet() and
  keyframes(): the storyboard as one picture, and a picture of every scene.
"""

from __future__ import annotations

import datetime as _dt
import difflib
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from PySide6.QtCore import QRectF, Qt
from PySide6.QtGui import QColor, QFont, QImage, QPainter

from . import canvas, graphics, media, scene_render, subs, templates
from .media import format_time
from .scenes import MIN_DURATION, TAIL, Project, Scene, TextBox, TextStyle

# field -> words its column header may be (spaces and marks are ignored when comparing)
FIELDS = {
    "number": ("장면", "번호", "장면번호", "no", "#", "씬", "scene", "컷", "cut", "순서", "순번"),
    "narration": ("나레이션", "내레이션", "대사", "음성", "스크립트", "대본", "narration", "script", "멘트", "voice"),
    "visual": ("화면", "화면설명", "설명", "연출", "비주얼", "visual", "장면설명", "내용", "그림설명", "화면구성"),
    "text": ("화면글자", "화면텍스트", "자막", "텍스트", "글자", "text", "caption", "타이틀", "제목"),
    "source": ("소스", "자료", "파일", "이미지", "영상", "source", "file", "그림", "소재"),
    "template": ("템플릿", "template", "틀", "레이아웃", "layout"),
    "params": ("값", "템플릿값", "params", "parameters", "설정", "변수"),
    "duration": ("길이", "시간", "초", "duration", "sec", "seconds", "길이초"),
    "note": ("메모", "비고", "note", "notes", "참고"),
    "credit": ("출처", "credit", "저작권", "라이선스", "license"),
}
LABELS = {"number": "장면", "narration": "나레이션", "visual": "화면 설명", "text": "화면 글자", "source": "소스",
          "template": "템플릿", "params": "템플릿 값", "duration": "길이(초)", "note": "메모", "credit": "출처"}  # fmt: skip
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="B7C3D0")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
THUMB_HEIGHT = 120  # pixels of a thumbnail in the workbook
READING_RATE = 6.5  # Korean characters a second, to guess a scene's length from its narration when there is no voice
PAUSE_ROW = 2.0  # seconds of silence for a row without narration, when none is given


def _norm(text) -> str:
    return re.sub(r"[\s_\-()\[\]·./:,]", "", str(text or "")).lower()


@dataclass
class Row:
    line: int  # its row in the sheet
    number: str = ""
    narration: str = ""
    visual: str = ""
    text: str = ""
    source: str = ""
    template: str = ""
    params: str = ""
    duration: float | None = None
    note: str = ""
    credit: str = ""


@dataclass
class Board:
    file: Path
    sheet: str
    header_row: int
    columns: dict[str, tuple[str, str]]  # field -> (column letter, its header)
    rows: list[Row] = field(default_factory=list)


# ---- reading --------------------------------------------------------------------------------------------


def _match(headers: dict[int, str]) -> dict[str, int]:
    """Columns for the fields: whole-word matches first, then headers that contain a word."""
    candidates = []
    order = list(FIELDS)
    for field_name, words in FIELDS.items():
        normal = [_norm(w) for w in words]
        for column, header in headers.items():
            if header in normal:
                candidates.append((2, -order.index(field_name), field_name, column))
            elif any(len(w) >= 2 and w in header for w in normal):
                candidates.append((1, -order.index(field_name), field_name, column))
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for _score, _priority, field_name, column in sorted(candidates, reverse=True):
        if field_name not in mapping and column not in used:
            mapping[field_name] = column
            used.add(column)
    return mapping


def _seconds(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, _dt.time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    if isinstance(value, _dt.timedelta):
        return value.total_seconds()
    text = str(value).strip().replace("초", "").replace("s", "")
    if ":" in text:
        parts = [float(p) for p in text.split(":")]
        total = 0.0
        for p in parts:
            total = total * 60 + p
        return total
    try:
        return float(text)
    except ValueError:
        return None


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, CellRichText):
        return str(value)
    return str(value).strip()


def read(path: str | Path, sheet: str = "", columns: dict | None = None) -> Board:
    """The rows of a 콘티 sheet. `columns` names columns by field ({"narration": "C"} or the header's words)."""
    path = Path(path)
    try:
        book = load_workbook(path, data_only=True, rich_text=False)
    except (OSError, ValueError, KeyError) as e:
        raise ValueError(f"엑셀 파일을 열 수 없습니다: {path.name}\n{e}") from e
    if sheet and sheet not in book.sheetnames:
        raise ValueError(f"'{sheet}' 시트가 없습니다. 있는 시트: {', '.join(book.sheetnames)}")
    ws = book[sheet] if sheet else next((book[n] for n in book.sheetnames if "콘티" in n), book.worksheets[0])
    best_row, best = 0, {}
    for r in range(1, min(ws.max_row, 15) + 1):
        headers = {c: _norm(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 60) + 1) if ws.cell(r, c).value not in (None, "")}
        found = _match(headers)
        if len(found) > len(best):
            best_row, best = r, found
    for field_name, where in (columns or {}).items():
        if field_name not in FIELDS:
            raise ValueError(f"columns의 '{field_name}'는 없는 항목입니다. 쓸 수 있는 것: {', '.join(FIELDS)}")
        where = str(where).strip()
        if re.fullmatch(r"[A-Za-z]{1,3}", where):
            best[field_name] = column_index_from_string(where.upper())
        else:
            hit = next((c for c in range(1, ws.max_column + 1) if _norm(ws.cell(best_row or 1, c).value) == _norm(where)), None)
            if hit is None:
                raise ValueError(f"'{where}' 머리글을 찾을 수 없습니다 ({field_name}).")
            best[field_name] = hit
    if not best_row or "narration" not in best and "source" not in best and "template" not in best:
        raise ValueError(f"{ws.title} 시트에서 머리글 줄(나레이션·화면설명·소스 등)을 찾지 못했습니다. columns로 열을 알려 주세요.")
    board = Board(path, ws.title, best_row, {f: (get_column_letter(c), _cell(ws.cell(best_row, c).value)) for f, c in best.items()})
    for r in range(best_row + 1, ws.max_row + 1):
        values = {f: ws.cell(r, c).value for f, c in best.items()}
        row = Row(r, **{f: _cell(values.get(f)) for f in ("number", "narration", "visual", "text", "source", "template",
                                                            "params", "note", "credit")})  # fmt: skip
        row.duration = _seconds(values.get("duration"))
        if any((row.narration, row.visual, row.text, row.source, row.template)):
            board.rows.append(row)
    return board


def describe(board: Board) -> str:
    lines = [f"{board.file.name} · 시트 '{board.sheet}' · 머리글 {board.header_row}번째 줄 · 장면 {len(board.rows)}개",
             "알아본 열: " + ", ".join(f"{LABELS[f]}={letter}({header})" for f, (letter, header) in board.columns.items())]  # fmt: skip
    missing = [LABELS[f] for f in ("narration", "visual", "source", "template", "duration") if f not in board.columns]
    if missing:
        lines.append(f"없는 열: {', '.join(missing)}")
    for k, row in enumerate(board.rows, 1):
        bits = [f"{k}. ({row.line}행{', ' + row.number if row.number else ''})"]
        for label, value in (("화면", row.visual), ("글자", row.text), ("소스", row.source), ("템플릿", row.template),
                             ("값", row.params), ("길이", f"{row.duration:g}초" if row.duration else ""), ("나레이션", row.narration)):  # fmt: skip
            if value:
                bits.append(f"{label}: {' '.join(str(value).split())[:70]}")
        lines.append(" · ".join(bits))
    return "\n".join(lines)


# ---- to a voice and a project ------------------------------------------------------------------------------


def script(board: Board) -> str:
    """The narration to make the voice from: a paragraph a row, a row without narration a pause."""
    parts = []
    for row in board.rows:
        text = " ".join(row.narration.split())
        parts.append(text if text else f"[쉼 {row.duration or PAUSE_ROW:g}초]")
    return "\n\n".join(parts)


def _plain(text: str) -> str:
    return re.sub(r"[\s\W_]+", "", text or "").lower()


def align(narrations: list[str], cues: list[subs.Cue]) -> tuple[list[tuple[int, int] | None], float]:
    """Which subtitle lines (from, to) each narration was spoken in, and how closely they match (0 to 1)."""
    spoken = [_plain(c.text) for c in cues]
    ranges: list[tuple[int, int] | None] = []
    ratios = []
    k = 0
    for text in narrations:
        target = _plain(text)
        if not target or k >= len(cues):
            ranges.append(None)
            continue
        best_end, best_ratio, acc = k + 1, -1.0, ""
        for end in range(k + 1, min(len(cues), k + 60) + 1):
            acc += spoken[end - 1]
            ratio = difflib.SequenceMatcher(None, target, acc, autojunk=False).ratio()
            if ratio > best_ratio:
                best_end, best_ratio = end, ratio
            if len(acc) > 1.6 * len(target) + 10:
                break
        ranges.append((k, best_end))
        ratios.append(best_ratio)
        k = best_end
    return ranges, (sum(ratios) / len(ratios) if ratios else 0.0)


def timings(rows: list[Row], cues: list[subs.Cue], voice_length: float) -> tuple[list[float], float]:
    """Each row's length, so the scenes change where the voice says each row's narration."""
    ranges, quality = align([r.narration for r in rows], cues)
    starts: list[float | None] = []
    for k, rng in enumerate(ranges):
        if rng is not None:
            starts.append(cues[rng[0]].start if k else 0.0)
        else:
            previous = next((ranges[j] for j in range(k - 1, -1, -1) if ranges[j] is not None), None)
            starts.append(cues[previous[1] - 1].end + 0.1 if previous else (0.0 if k == 0 else None))
    for k in range(len(starts)):  # a row nothing pins down sits right after the one before it
        if starts[k] is None:
            starts[k] = (starts[k - 1] or 0.0) + (rows[k - 1].duration or PAUSE_ROW)
        if k and starts[k] < starts[k - 1] + MIN_DURATION:
            starts[k] = starts[k - 1] + MIN_DURATION
    end = max(voice_length, cues[-1].end if cues else 0.0) + TAIL
    lengths = [max(MIN_DURATION, (starts[k + 1] if k + 1 < len(starts) else end) - starts[k]) for k in range(len(starts))]
    return lengths, quality


def kind_of(path: Path) -> str | None:
    suffix = path.suffix.lower()
    if suffix in canvas.IMAGE_SUFFIXES:
        return "image"
    if suffix in media.VIDEO_SUFFIXES:
        return "clip"
    return None


def auto_fit(project: Project, path: Path, kind: str, tolerance: float = 1.43) -> str:
    """"cover" when the picture has about the frame's shape, otherwise "contain" (shown whole)."""
    width, height = project.size
    try:
        if kind == "image":
            source_w, source_h = canvas.sticker_size(str(path))
        else:
            info = media.probe(path)
            source_w, source_h = info.width, info.height
    except (OSError, ValueError):
        return "cover"
    if not source_w or not source_h:
        return "cover"
    ratio = (source_w / source_h) / (width / height)
    return "cover" if 1 / tolerance <= ratio <= tolerance else "contain"


def _source(row: Row, folder: Path) -> tuple[Path | None, str]:
    """The row's picture or clip (relative to the workbook's folder), or why there is none."""
    for name in re.split(r"[\n|;]", row.source):
        name = name.strip().strip('"')
        if not name:
            continue
        path = Path(name)
        for candidate in ([path] if path.is_absolute() else [folder / path, folder / "소스" / path]):
            if candidate.is_file() and kind_of(candidate):
                return candidate, ""
        return None, f"파일을 찾을 수 없습니다: {name}"
    return None, ""


def _estimate(narration: str) -> float:
    return max(MIN_DURATION, len(_plain(narration)) / READING_RATE + 0.6)


def build(board: Board, base: Project | None = None, voice: Path | None = None) -> tuple[Project, list[str]]:
    """A project with a scene a row (settings from `base`), timed to `voice` when given, and notes on what it did."""
    project = Project.from_json(base.to_json()) if base else Project()
    project.scenes, project.cues = [], None
    notes: list[str] = []
    folder = board.file.parent
    for k, row in enumerate(board.rows, 1):
        path, problem = _source(row, folder)
        if problem:
            notes.append(f"{k}번 장면({row.line}행): {problem}")
        scene = None
        if row.template:
            try:
                template = templates.get(row.template)
                values = templates.parse_params(row.params, template)
                blanks = templates.params(template)
                for blank in ("title", "text", "heading"):  # on-screen words fill the template's first word blank
                    if row.text and blank in blanks and blank not in values:
                        values[blank] = row.text
                        break
                if path is not None and "image" in blanks and "image" not in values:
                    values["image"] = str(path)
                filled = templates.expand(template, values)
                problems = graphics.check_layers(filled["layers"], project.style)
                if problems:
                    raise ValueError(problems[0])
                scene = Scene(kind="graphic", style=project.style, layers=filled["layers"], duration=filled["duration"])
            except ValueError as e:
                notes.append(f"{k}번 장면({row.line}행): 템플릿 '{row.template}'을(를) 쓰지 못했습니다 — {e}")
        if scene is None and path is not None:
            kind = kind_of(path)
            scene = Scene(kind=kind, source=str(path), fit=auto_fit(project, path, kind))
            if kind == "clip":
                length = media.probe(path).video_duration
                scene.duration = max(MIN_DURATION, length) if length else scene.duration
            if row.text:
                scene.texts = [TextBox(row.text, TextStyle(size=64), y=0.85)]
        if scene is None and row.text:
            filled = templates.expand("제목", {"title": row.text})
            scene = Scene(kind="graphic", style=project.style, layers=filled["layers"], duration=filled["duration"])
        if scene is None:
            scene = Scene()  # an empty scene, to be filled
        scene.narration = " ".join(row.narration.split())
        if row.duration:
            scene.duration = max(MIN_DURATION, row.duration)
        elif row.narration and not voice:
            scene.duration = _estimate(row.narration)
        project.scenes.append(scene)
    if voice is not None:
        srt = next((voice.with_suffix(s) for s in subs.SUBTITLE_SUFFIXES if voice.with_suffix(s).exists()), None)
        if srt is None:
            notes.append("음성 옆에 같은 이름의 자막(.srt)이 없어서 길이를 음성에 맞추지 못했습니다.")
        else:
            cues = subs.load_cues(srt)
            lengths, quality = timings(board.rows, cues, media.probe(voice).duration)
            for scene, length in zip(project.scenes, lengths):
                scene.duration = length
            project.voice = str(voice)
            notes.append(f"장면 길이를 음성에 맞췄습니다 (콘티 나레이션과 음성 자막 일치 {quality * 100:.0f}%).")
            if quality < 0.85:
                notes.append("일치율이 낮습니다. 음성이 이 콘티의 대본으로 만든 것인지 compare_narration으로 확인하세요.")
    return project, notes


# ---- pictures of scenes -------------------------------------------------------------------------------------


def moment(scene: Scene) -> float:
    """The moment that shows a scene best: a graphic scene at its fullest, others a little way in."""
    if scene.kind == "graphic":
        return graphics.poster_time(scene.layers, scene.style, scene.duration)
    return min(scene.duration * 0.4, max(0.0, scene.duration - 0.05))


def picture(project: Project, index: int, width: int | None = None) -> QImage:
    frame_w, frame_h = project.size
    image = scene_render.still(project.scenes[index], frame_w, frame_h, moment(project.scenes[index]),
                               canvas.logo_layer(project, frame_w, frame_h))  # fmt: skip
    if width:
        image = image.scaledToWidth(width, Qt.TransformationMode.SmoothTransformation)
    return image


def _starts(project: Project) -> list[float]:
    out, t = [], 0.0
    for scene in project.scenes:
        out.append(t)
        t += scene.duration
    return out


def contact_sheet(project: Project, title: str = "", columns: int = 4, width: int = 2400) -> QImage:
    """The storyboard on one picture: every scene's picture with its number, time and narration under it."""
    columns = max(1, min(8, columns))
    frame_w, frame_h = project.size
    cell_w = width // columns
    pad = round(cell_w * 0.04)
    thumb_w = cell_w - 2 * pad
    thumb_h = round(thumb_w * frame_h / frame_w)
    text_h = round(cell_w * 0.2)
    cell_h = thumb_h + text_h + pad
    head = round(width * 0.04)
    rows = -(-len(project.scenes) // columns) if project.scenes else 1
    sheet = QImage(width, head + rows * cell_h + pad, QImage.Format.Format_RGB32)
    sheet.fill(QColor("#ffffff"))
    painter = QPainter(sheet)
    try:
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform)
        heading = QFont("Malgun Gothic")
        heading.setPixelSize(max(12, round(head * 0.5)))
        heading.setBold(True)
        painter.setFont(heading)
        painter.setPen(QColor("#111827"))
        painter.drawText(QRectF(pad, 0, width - 2 * pad, head), Qt.AlignmentFlag.AlignVCenter,
                         f"{title or '스토리보드'} · 장면 {len(project.scenes)}개 · {format_time(project.duration, precise=True)}")  # fmt: skip
        small, bold = QFont("Malgun Gothic"), QFont("Malgun Gothic")
        small.setPixelSize(max(10, round(text_h * 0.17)))
        bold.setPixelSize(max(10, round(text_h * 0.2)))
        bold.setBold(True)
        for k, (scene, start) in enumerate(zip(project.scenes, _starts(project))):
            x = (k % columns) * cell_w + pad
            y = head + (k // columns) * cell_h
            painter.drawImage(QRectF(x, y, thumb_w, thumb_h), picture(project, k, thumb_w))
            painter.setPen(QColor("#d1d5db"))
            painter.drawRect(QRectF(x, y, thumb_w, thumb_h))
            painter.setPen(QColor("#111827"))
            painter.setFont(bold)
            line = f"{k + 1}. {format_time(start, precise=True)}–{format_time(start + scene.duration, precise=True)} ({scene.duration:.1f}초)"
            painter.drawText(QRectF(x, y + thumb_h + pad * 0.3, thumb_w, text_h * 0.3), Qt.AlignmentFlag.AlignLeft, line)
            painter.setFont(small)
            painter.setPen(QColor("#374151"))
            words = scene.narration or ("(나레이션 없음)" if not scene.text else scene.text)
            painter.drawText(QRectF(x, y + thumb_h + pad * 0.3 + text_h * 0.32, thumb_w, text_h * 0.66),
                             Qt.TextFlag.TextWordWrap | Qt.AlignmentFlag.AlignLeft, words)  # fmt: skip
    finally:
        painter.end()
    return sheet


def keyframes(project: Project, folder: Path) -> list[Path]:
    """Every scene's picture at full size, as 장면_01.png ..."""
    folder.mkdir(parents=True, exist_ok=True)
    out = []
    for k in range(len(project.scenes)):
        path = folder / f"장면_{k + 1:02d}.png"
        if not picture(project, k).save(str(path)):
            raise OSError(f"그림을 저장하지 못했습니다: {path}")
        out.append(path)
    return out


# ---- the narration check -----------------------------------------------------------------------------------


@dataclass
class Check:
    number: int
    written: str
    spoken: str
    ratio: float


def compare(narrations: list[str], cues: list[subs.Cue]) -> list[Check]:
    """Each narration against what the voice says in its place (which subtitle lines, found by the words)."""
    ranges, _quality = align(narrations, cues)
    out = []
    for k, (text, rng) in enumerate(zip(narrations, ranges), 1):
        spoken = " ".join(c.text for c in cues[rng[0] : rng[1]]) if rng else ""
        written = " ".join(text.split())
        a, b = _plain(written), _plain(spoken)
        ratio = 1.0 if a == b else difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()
        if written or spoken:
            out.append(Check(k, written, " ".join(spoken.split()), ratio))
    return out


def _rich(a: str, b: str) -> tuple[CellRichText, CellRichText]:
    """Both texts with what differs marked: missing from the voice in red, said differently in blue."""
    left, right = CellRichText(), CellRichText()
    gone = InlineFont(color="C00000", b=True, strike=True)
    added = InlineFont(color="1F4E9F", b=True, u="single")
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == "equal":
            left.append(a[i1:i2])
            right.append(b[j1:j2])
            continue
        if i2 > i1:
            left.append(TextBlock(gone, a[i1:i2]))
        if j2 > j1:
            right.append(TextBlock(added, b[j1:j2]))
    return left, right


def _header(ws, labels: list[str], widths: list[float]) -> None:
    for c, (label, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(1, c, label)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


def write_check(book: Workbook, checks: list[Check], title: str = "나레이션 비교") -> None:
    ws = book.create_sheet(title)
    _header(ws, ["장면", "일치", "콘티 나레이션 (빨강: 음성에 없음)", "음성 자막 (파랑: 다르게 말함)"], [7, 9, 60, 60])
    for r, check in enumerate(checks, 2):
        left, right = _rich(check.written, check.spoken)
        ws.cell(r, 1, check.number)
        ratio = ws.cell(r, 2, round(check.ratio, 3))
        ratio.number_format = "0%"
        if check.ratio < 0.98:
            ratio.fill = PatternFill("solid", fgColor="FDE2E1")
        ws.cell(r, 3).value = left if len(left) else check.written
        ws.cell(r, 4).value = right if len(right) else check.spoken
        for c in range(1, 5):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, c).border = GRID


# ---- the workbook -------------------------------------------------------------------------------------------


def _scene_kind(scene: Scene) -> str:
    if scene.kind == "graphic":
        return "그래픽"
    if scene.kind in ("image", "clip") and scene.source:
        return "이미지" if scene.kind == "image" else "영상"
    return "빈 장면" if scene.is_blank else "색 화면"


def _words(scene: Scene) -> str:
    words = [t.text.strip() for t in scene.texts if t.text.strip()]

    def walk(layers):
        for layer in layers:
            if isinstance(layer, dict):
                if layer.get("type") == "text" and layer.get("text"):
                    words.append(re.sub(r"\[/?[^\[\]]*\]", "", str(layer["text"])))
                walk(layer.get("layers") or [])

    walk(scene.layers)
    return "\n".join(words)


def sources(project: Project) -> list[dict]:
    """Every file the video uses: what it is, where, and in which scenes."""
    found: dict[str, dict] = {}

    def add(path: str, kind: str, scene: int | None = None) -> None:
        if not path:
            return
        entry = found.setdefault(str(Path(path)), {"kind": kind, "path": str(Path(path)), "scenes": []})
        if scene and scene not in entry["scenes"]:
            entry["scenes"].append(scene)

    for k, scene in enumerate(project.scenes, 1):
        if scene.kind in ("image", "clip") and scene.source:
            add(scene.source, "이미지" if scene.kind == "image" else "영상", k)
        for overlay in scene.overlays:
            add(overlay.source, "겹친 그림" if overlay.kind == "image" else "겹친 영상", k)
        if scene.kind == "graphic":
            for path in graphics.files(scene.layers, scene.style):
                add(path, "그래픽 재료", k)
    if project.voice:
        add(project.voice, "음성")
    if project.music:
        add(project.music, "배경 음악")
    if project.logo:
        add(project.logo, "로고")
    for sound in project.sounds:
        add(sound.source, "효과음")
    return list(found.values())


def export(project: Project, output: Path, thumbnails: bool = True, title: str = "") -> list[str]:
    """The project as a 콘티 workbook: a row a scene (with its picture), the source list, and the narration check."""
    book = Workbook()
    ws = book.active
    ws.title = "콘티"
    labels = ["장면", "시작", "끝", "길이(초)", "화면", "나레이션", "종류", "소스", "템플릿 값", "화면 글자", "메모"]
    frame_w, frame_h = project.size
    thumb_w = round(THUMB_HEIGHT * frame_w / frame_h)
    _header(ws, labels, [7, 10, 10, 9, thumb_w / 7 + 2, 50, 9, 28, 24, 28, 20])
    notes = []
    with tempfile.TemporaryDirectory(prefix="movie_board_") as tmp:
        for r, (scene, start) in enumerate(zip(project.scenes, _starts(project)), 2):
            used = [Path(scene.source).name] if scene.kind in ("image", "clip") and scene.source else []
            used += [Path(o.source).name for o in scene.overlays if o.source]
            values = [r - 1, format_time(start, precise=True), format_time(start + scene.duration, precise=True), round(scene.duration, 2),
                      "", scene.narration, _scene_kind(scene), "\n".join(used),
                      ", ".join(dict.fromkeys(str(layer.get("type")) for layer in scene.layers)) if scene.kind == "graphic" else "",
                      _words(scene), ""]  # fmt: skip
            for c, value in enumerate(values, 1):
                cell = ws.cell(r, c, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = GRID
            if thumbnails:
                try:
                    file = Path(tmp) / f"t{r}.png"
                    picture(project, r - 2, thumb_w).save(str(file))
                    image = SheetImage(str(file))
                    ws.add_image(image, f"E{r}")
                    ws.row_dimensions[r].height = THUMB_HEIGHT * 0.75 + 4  # points
                except (OSError, ValueError, RuntimeError) as e:
                    notes.append(f"{r - 1}번 장면 그림을 넣지 못했습니다: {e}")
        list_ws = book.create_sheet("소스목록")
        _header(list_ws, ["번호", "종류", "파일 이름", "경로", "쓰인 장면", "출처", "라이선스", "메모"], [6, 11, 28, 60, 14, 24, 16, 20])
        for r, entry in enumerate(sources(project), 2):
            values = [r - 1, entry["kind"], Path(entry["path"]).name, entry["path"], ", ".join(map(str, entry["scenes"])), "", "", ""]
            for c, value in enumerate(values, 1):
                cell = list_ws.cell(r, c, value)
                cell.border = GRID
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        voice = Path(project.voice) if project.voice else None
        cues = []
        if project.cues is not None:
            cues = project.cues
        elif voice is not None:
            srt = next((voice.with_suffix(s) for s in subs.SUBTITLE_SUFFIXES if voice.with_suffix(s).exists()), None)
            cues = subs.load_cues(srt) if srt else []
        if cues and any(s.narration for s in project.scenes):
            write_check(book, compare([s.narration for s in project.scenes], cues))
        output.parent.mkdir(parents=True, exist_ok=True)
        book.save(str(output))
    return notes


def blank(output: Path) -> None:
    """An empty 콘티 workbook in the layout read() knows best, with a sheet explaining the columns and templates."""
    book = Workbook()
    ws = book.active
    ws.title = "콘티"
    labels = ["장면", "나레이션", "화면 설명", "화면 글자", "소스", "템플릿", "템플릿 값", "길이(초)", "메모"]
    _header(ws, labels, [7, 50, 30, 24, 26, 12, 34, 9, 20])
    examples = [
        [1, "", "제목 화면", "외팔보의 굽힘 응력", "", "제목", "subtitle=하중이 끝에 걸릴 때", 3, "나레이션이 없으면 그 길이만큼 쉼"],
        [2, "외팔보 끝에 하중이 걸리면 고정단 쪽에 큰 응력이 생깁니다.", "해석 결과 색 타일", "", "", "해석결과", "heading=굽힘 응력; max=182", "", ""],
        [3, "최대 응력은 백팔십이 메가파스칼입니다.", "숫자 강조", "", "", "큰숫자", 'value=182; unit=" MPa"; label=최대 응력', "", ""],
        [4, "현장 사진으로 확인해 보겠습니다.", "현장 사진", "", "사진1.jpg", "", "", "", "소스는 이 엑셀과 같은 폴더(또는 소스 폴더) 기준"],
    ]
    for r, values in enumerate(examples, 2):
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = GRID
    guide = book.create_sheet("도움말")
    _header(guide, ["항목", "설명"], [16, 110])
    rows = [
        ("장면", "번호 (없어도 됨)"),
        ("나레이션", "이 장면에서 읽을 말. 대본(storyboard_script)으로 음성을 만들면, 장면 길이가 음성에 맞춰집니다."),
        ("화면 설명", "어떤 화면인지 적어 두는 칸 (만들 때는 쓰지 않음)"),
        ("화면 글자", "화면에 띄울 글자. 소스가 있으면 그 위 아래쪽에, 없으면 제목 화면으로. 템플릿이 있으면 템플릿의 제목 칸에"),
        ("소스", "이미지·영상 파일 이름이나 경로 (이 엑셀이 있는 폴더, 또는 그 안의 '소스' 폴더 기준)"),
        ("템플릿", "그래픽 장면 틀 이름: " + ", ".join(templates.names())),
        ("템플릿 값", "이름=값; 이름=값 (목록은 가 | 나 | 다, 앞뒤 빈칸을 살리려면 \" MPa\"처럼 따옴표). 템플릿마다 받는 값은 아래"),
        ("길이(초)", "장면 길이. 음성에 맞추면 무시되고, 나레이션 없는 줄에서는 쉬는 시간"),
    ]
    for name in templates.names():
        spec = templates.params(templates.get(name))
        rows.append((f"템플릿 {name}", templates.get(name).get("description", "") + "\n" +
                     "\n".join(f"  {p}: {s.get('doc', '')} (기본값 {templates.as_text(s.get('default'))})" for p, s in spec.items())))  # fmt: skip
    for r, (name, text) in enumerate(rows, 2):
        guide.cell(r, 1, name).alignment = Alignment(vertical="top")
        guide.cell(r, 2, text).alignment = Alignment(wrap_text=True, vertical="top")
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(output))
